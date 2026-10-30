"""把 attachments 目录里的附件全部提取成文本，供条目流水线使用。

- PDF: pdftotext（布局模式），失败退回 pypdf
- DOCX: python-docx（含表格）
- DOC: antiword（UTF-8 映射）
- XLSX: openpyxl；XLS: xlrd
- RAR: WinRAR UnRAR 解包后逐个提取
名单类文件含学生姓名，不提取、不入库。

用法：python tools/parse_attachments_20261030.py [--limit N]
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "sources"
ATT = SRC / "attachments"
OUT = ATT / "text"
UNRAR = Path(r"C:\Program Files\WinRAR\UnRAR.exe")

PRIVACY = re.compile(r"录取名单|拟录取|获奖名单|名单公示|学生名单|复试名单|拟推荐名单")
INNER_EXT = {"pdf", "doc", "docx", "xls", "xlsx"}


def sh(args: list[str], timeout: int = 180) -> str:
    r = subprocess.run(args, capture_output=True, timeout=timeout)
    if r.returncode != 0:
        raise RuntimeError((r.stderr or b"").decode("utf-8", "ignore")[:200])
    return r.stdout.decode("utf-8", "ignore")


def pdf_text(p: Path) -> str:
    try:
        return sh(["pdftotext", "-layout", "-enc", "UTF-8", str(p), "-"])
    except Exception:
        from pypdf import PdfReader

        return "\n".join((pg.extract_text() or "") for pg in PdfReader(str(p)).pages)


def docx_text(p: Path) -> str:
    from docx import Document

    d = Document(str(p))
    parts = [x.text for x in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return "\n".join(parts)


def doc_text(p: Path) -> str:
    return sh(["antiword", "-m", "UTF-8.txt", str(p)])


def xlsx_text(p: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(p), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"[{ws.title}]")
        for row in ws.iter_rows(values_only=True):
            vals = [str(v).strip() for v in row if v not in (None, "")]
            if vals:
                parts.append(" | ".join(vals))
    wb.close()
    return "\n".join(parts)


def xls_text(p: Path) -> str:
    import xlrd

    wb = xlrd.open_workbook(str(p))
    parts = []
    for ws in wb.sheets():
        parts.append(f"[{ws.name}]")
        for r in range(ws.nrows):
            vals = [str(v).strip() for v in ws.row_values(r) if str(v).strip()]
            if vals:
                parts.append(" | ".join(vals))
    return "\n".join(parts)


def extract(p: Path) -> str:
    ext = p.suffix.lower().lstrip(".")
    if ext == "pdf":
        return pdf_text(p)
    if ext == "docx":
        return docx_text(p)
    if ext == "doc":
        return doc_text(p)
    if ext == "xlsx":
        return xlsx_text(p)
    if ext == "xls":
        return xls_text(p)
    raise RuntimeError(f"不支持的格式 {ext}")


def rar_text(p: Path) -> str:
    if not UNRAR.exists():
        raise RuntimeError("找不到 UnRAR.exe")
    with tempfile.TemporaryDirectory() as td:
        subprocess.run([str(UNRAR), "x", "-o+", "-inul", str(p), td + "\\"],
                       capture_output=True, timeout=300)
        parts = []
        for f in sorted(Path(td).rglob("*")):
            if f.is_file() and f.suffix.lower().lstrip(".") in INNER_EXT:
                try:
                    parts.append(f"[{f.name}]\n{extract(f)}")
                except Exception as e:  # noqa: BLE001
                    parts.append(f"[{f.name}] 解析失败：{e}")
        return "\n\n".join(parts)


def norm(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\u3000]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--limit", type=int, default=0)
    args = ap.parse_args()

    idx = json.loads((ATT.parent / "attachments-index.json").read_text(encoding="utf-8"))
    OUT.mkdir(parents=True, exist_ok=True)
    out_index: dict[str, dict] = {}
    stat = {"ok": 0, "short": 0, "privacy_skip": 0, "error": 0}

    rows = sorted(idx.values(), key=lambda v: v["page_title"])
    if args.limit:
        rows = rows[: args.limit]
    for v in rows:
        slug = "att-" + hashlib.sha1(v["url"].encode()).hexdigest()[:10]
        name = v["name"]
        title = f"{v['page_title']} · {re.sub(r'\.(pdf|docx?|xlsx?)$', '', name, flags=re.I)}"
        if PRIVACY.search(v["page_title"]) or PRIVACY.search(name):
            stat["privacy_skip"] += 1
            print(f"  [私] {title[:60]}")
            continue
        text_p = OUT / f"{slug}.txt"
        if text_p.exists() and text_p.stat().st_size > 50:
            text = text_p.read_text(encoding="utf-8")
        else:
            p = ROOT / v["saved"]
            try:
                text = norm(rar_text(p) if v["ext"] == "rar" else extract(p))
            except Exception as e:  # noqa: BLE001
                stat["error"] += 1
                print(f"  ✗ {title[:56]}：{e}")
                continue
            text_p.write_text(text, encoding="utf-8")
        chars = len(text)
        status = "ok" if chars >= 200 else "short"
        stat[status] += 1
        out_index[slug] = {
            "url": v["url"],
            "page_url": v["page_url"],
            "page_slug": v["page_slug"],
            "page_title": v["page_title"],
            "title": title,
            "name": name,
            "ext": v["ext"],
            "chars": chars,
            "bytes": v["bytes"],
            "source_sha256": v["sha256"],
            "kind": "attachment",
            "fetched_at": v["fetched_at"],
            "status": status,
        }
        print(f"  {'✓' if status=='ok' else '·'} {chars:>6} 字  {title[:60]}")

    (SRC / "attachments-text-index.json").write_text(
        json.dumps(out_index, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"\n完成：{stat}，文本索引 {len(out_index)} 条")


if __name__ == "__main__":
    main()
