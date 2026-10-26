"""把 cdp 语料里"正文只有附件"的页面，附件下载下来提取文本再入库。

用法：python tools/fetch_attachments.py --slugs 2025... ,python 任意
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
import time
import urllib.parse
import urllib.request
from datetime import datetime, timedelta, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
from crawl_sitemap import body_text  # noqa: E402

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/141.0.0.0 Safari/537.36")
TZ = timezone(timedelta(hours=8))
OUT = ROOT / "sources" / "cdp"
IDX = ROOT / "sources" / "cdp-index.json"


def get(url: str, timeout: int = 40) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read()


def attachments(html: str, base: str) -> list[tuple[str, str]]:
    out = []
    for href, text in re.findall(r'<a[^>]+href="([^"]+\.(?:pdf|docx?|xlsx?))"[^>]*>(.*?)</a>', html, re.S | re.I):
        name = re.sub(r"<[^>]+>", "", text).strip() or href.rsplit("/", 1)[-1]
        out.append((urllib.parse.urljoin(base, href), name))
    return out


def pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    return "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)


def docx_text(path: Path) -> str:
    from docx import Document

    d = Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def main() -> None:
    slugs = sys.argv[sys.argv.index("--slugs") + 1].split(",") if "--slugs" in sys.argv else []
    index = json.loads(IDX.read_text(encoding="utf-8"))
    todo = [s for s in slugs if s in index] if slugs else [
        s for s, m in index.items() if m.get("chars", 0) < 900 and "奖学金" in m.get("title", "")
    ]
    print(f"待处理页面 {len(todo)} 个")

    added = 0
    for slug in todo:
        meta = index[slug]
        try:
            html = get(meta["url"]).decode("utf-8", "ignore")
        except Exception as e:  # noqa: BLE001
            print(f"  ✗ 取页失败 {slug}: {e}")
            continue
        atts = attachments(html, meta["url"])
        print(f"\n{meta['title'][:40]}  附件 {len(atts)} 个")
        for url, name in atts[:3]:
            ext = url.rsplit(".", 1)[-1].lower()
            if ext not in ("pdf", "doc", "docx"):
                continue
            tmp = OUT / f"_tmp.{ext}"
            try:
                tmp.write_bytes(get(url, timeout=90))
            except Exception as e:  # noqa: BLE001
                print(f"    ✗ 下载失败 {name[:30]}: {e}")
                continue
            try:
                text = pdf_text(tmp) if ext == "pdf" else docx_text(tmp)
            except Exception as e:  # noqa: BLE001
                print(f"    ✗ 解析失败 {name[:30]}: {e}")
                continue
            finally:
                tmp.unlink(missing_ok=True)
            if len(text) < 300:
                print(f"    ✗ 文本过少（可能是扫描件） {name[:30]}")
                continue
            s2 = f"{slug}-att{len(name) % 10}"
            (OUT / f"{s2}.txt").write_text(text, encoding="utf-8")
            index[s2] = {
                "url": url, "title": f"{meta['title'][:28]}·附件 {name[:24]}", "chars": len(text),
                "sha256": hashlib.sha256(text.encode()).hexdigest()[:16],
                "published": meta.get("published", ""),
                "fetched_at": datetime.now(TZ).isoformat(timespec="seconds"),
                "kind": "cdp-att", "host": meta.get("host", ""),
            }
            added += 1
            print(f"    [收] {len(text):>5} 字  {name[:34]}")
            time.sleep(0.3)

    IDX.write_text(json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\n新增附件文本 {added} 条，索引 {len(index)} 条")


if __name__ == "__main__":
    main()
