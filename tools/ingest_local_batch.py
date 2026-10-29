"""把用户下载的本地文件（PDF/DOCX/XLSX）提取成文本，入 sources/local。

用法：python tools/ingest_local_batch.py --dir "<素材目录>"
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
from datetime import datetime, timedelta, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "sources" / "local"
KEEP = ROOT / "sources" / "local" / "_files"
TZ = timezone(timedelta(hours=8))


def pdf_text(p: Path) -> str:
    from pypdf import PdfReader

    r = PdfReader(str(p))
    return "\n".join((pg.extract_text() or "") for pg in r.pages)


def docx_text(p: Path) -> str:
    from docx import Document

    d = Document(str(p))
    parts = [x.text for x in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def xlsx_text(p: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(p), read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"== {ws.title} ==")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def slug_of(name: str) -> str:
    base = re.sub(r"\.(pdf|docx?|xlsx?)$", "", name, flags=re.I)
    base = re.sub(r"[^\w\u4e00-\u9fff]+", "-", base).strip("-")
    return "loc-" + base[:40] + "-" + hashlib.md5(name.encode()).hexdigest()[:6]


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", required=True)
    args = ap.parse_args()

    src_dir = Path(args.dir)
    OUT.mkdir(parents=True, exist_ok=True)
    KEEP.mkdir(parents=True, exist_ok=True)
    idx_p = ROOT / "sources" / "local-index.json"
    index = json.loads(idx_p.read_text(encoding="utf-8")) if idx_p.exists() else {}

    files = [p for p in sorted(src_dir.iterdir()) if p.is_file()]
    print(f"发现 {len(files)} 个文件")

    added = 0
    for p in files:
        ext = p.suffix.lower()
        slug = slug_of(p.name)
        if slug in index:
            print(f"  已入库 {p.name[:30]}")
            continue
        try:
            if ext == ".pdf":
                text = pdf_text(p)
            elif ext in (".docx", ".doc"):
                text = docx_text(p)
            elif ext in (".xlsx", ".xls"):
                text = xlsx_text(p)
            else:
                print(f"  跳过（不支持）{p.name[:30]}")
                continue
        except Exception as e:  # noqa: BLE001
            print(f"  ✗ {p.name[:30]}: {e}")
            continue
        if len(text) < 120:
            print(f"  ✗ 文本过少 {p.name[:30]} ({len(text)} 字)")
            continue
        (OUT / f"{slug}.txt").write_text(text, encoding="utf-8")
        shutil.copy2(p, KEEP / p.name)
        index[slug] = {
            "url": "", "title": re.sub(r"\.(pdf|docx?|xlsx?)$", "", p.name, flags=re.I),
            "chars": len(text), "bytes": p.stat().st_size, "file": p.name,
            "sha256": hashlib.sha256(p.read_bytes()).hexdigest()[:16],
            "fetched_at": datetime.now(TZ).isoformat(timespec="seconds"),
            "kind": "local", "official": True,
        }
        added += 1
        print(f"  [收] {len(text):>6} 字  {p.name[:38]}")

    idx_p.write_text(json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\n新增 {added} 份，索引 {len(index)} 条")


if __name__ == "__main__":
    main()
